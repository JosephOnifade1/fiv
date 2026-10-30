"""Gig Builder — full Fiverr gig setups generated from scan data."""
import asyncio
import io
import re
import statistics
import uuid
from datetime import datetime, timezone

try:
    from docx import Document
    from docx.shared import Pt
    HAS_DOCX = True
except ImportError:
    Document = None  # type: ignore
    Pt = None  # type: ignore
    HAS_DOCX = False

from fastapi import APIRouter, BackgroundTasks, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from typing import Annotated, Any, Optional

from auth import get_current_user
from billing import is_pro
from scraper import fetch_autocomplete
import os
from ai_insights import _parse_json, run_chat

logger = __import__("logging").getLogger("fiverr-scout.gig_builder")

gig_router = APIRouter(prefix="/api/gigs")

SECTIONS = ["titles", "category", "tags", "pricing", "description", "faq", "requirements", "gallery"]
FREE_SECTIONS = {"titles", "tags", "category"}
LIMITS = {"title": 80, "tag": 20, "tags": 5, "description": 1200, "package_name": 35,
          "package_description": 100, "min_price": 5, "gallery": 3}
STALE_DAYS = 7
PRICE_STEPS = [5, 10, 15, 20, 25, 30, 35, 40, 45, 50, 60, 70, 75, 80, 90, 100, 120, 125, 150, 175,
               200, 225, 250, 300, 350, 400, 450, 500, 600, 750, 800, 1000, 1200, 1500, 2000, 2500, 3000]
OFF_PLATFORM = re.compile(
    r"\b(e-?mail|gmail|whats\s?app|skype|telegram|paypal|phone|call me|contact me outside|"
    r"discord|instagram|facebook|linkedin|www\.|https?://)\S*", re.I)
EMAILISH = re.compile(r"\S+@\S+\.\S+")


class BuildIn(BaseModel):
    job_id: Optional[str] = None
    keyword: Optional[str] = None


class RegenIn(BaseModel):
    section: str


class PatchIn(BaseModel):
    section: str
    value: Any


class OrigPackage(BaseModel):
    tier: str
    name: str = Field("", max_length=200)
    description: str = Field("", max_length=500)
    price: Optional[float] = Field(None, ge=0, le=100000)
    delivery_days: Optional[int] = Field(None, ge=0, le=365)


class OrigFaq(BaseModel):
    question: str = Field(..., max_length=500)
    answer: str = Field(..., max_length=2000)


class OriginalGig(BaseModel):
    title: str = Field(..., min_length=5, max_length=200)
    tags: list[Annotated[str, Field(max_length=60)]] = Field(default_factory=list, max_length=10)
    description: str = Field("", max_length=5000)
    packages: list[OrigPackage] = Field(default_factory=list, max_length=3)
    faq: list[OrigFaq] = Field(default_factory=list, max_length=15)


class AuditIn(BaseModel):
    job_id: str
    original: OriginalGig


GIG_SYSTEM = (
    "You are a top Fiverr SEO copywriter who writes gigs that rank on page one. You write plain, "
    "buyer-focused English that matches real search phrases. Never invent futuristic buzzwords or trends "
    "that are not in the provided data. Output STRICT JSON only."
)


def _gig_chat():
    """Return a chat object if an LLM key is available, else None."""
    key = os.environ.get("GEMINI_API_KEY") or os.environ.get("EMERGENT_LLM_KEY", "")
    if not key:
        return None
    try:
        from emergentintegrations.llm.chat import LlmChat
        return LlmChat(api_key=key, session_id=f"gig-{uuid.uuid4()}",
                       system_message=GIG_SYSTEM).with_model("gemini", "gemini-3-flash-preview")
    except Exception:
        return None


# ---------------- pricing anchors ----------------
def _snap(x: float) -> int:
    return min(PRICE_STEPS, key=lambda s: abs(s - x))


def _next_step(p: int) -> int:
    return next((s for s in PRICE_STEPS if s > p), p + 50)


def price_anchors(gigs: list[dict], metrics: dict) -> dict:
    prices = sorted(g["price_from"] for g in gigs if g.get("price_from"))
    median = metrics.get("median_price") or (statistics.median(prices) if prices else 25)
    top = metrics.get("top_decile_price") or (prices[int(len(prices) * 0.9)] if prices else median * 3)
    below = [s for s in PRICE_STEPS if s < median]
    basic = below[-1] if below else 5
    standard = _snap(median)
    if standard <= basic:
        standard = _next_step(basic)
    premium = _snap(max(top, standard * 1.8))
    if premium <= standard:
        premium = _next_step(standard)
    return {"median": round(median, 2), "top_decile": round(top, 2),
            "basic": basic, "standard": standard, "premium": premium}


# ---------------- prompts ----------------
def _ctx_block(ctx: dict) -> str:
    original = ""
    if ctx.get("original"):
        original = ("\nSELLER'S EXISTING GIG (reference only — reuse their real services, deliverables and strengths, "
                    "but answer ONLY with the exact JSON asked below, nothing else):\n"
                    f"{ctx['original']}\n")
    return (
        f"MAIN KEYWORD: \"{ctx['keyword']}\"\n"
        f"Top-ranking gig titles: {ctx['top_titles']}\n"
        f"Tags used by top gigs: {ctx['tags']}\n"
        f"Fiverr autocomplete phrases (high search demand): {ctx['autocomplete']}\n"
        f"Price stats USD: median {ctx['anchors']['median']}, top decile {ctx['anchors']['top_decile']}\n"
        f"Typical delivery days of top gigs: {ctx['delivery_days']}\n"
        f"{original}"
    )


AUDIT_PROMPT = (
    "Audit the SELLER'S EXISTING GIG against the niche data above. List 5-8 concrete problems that hurt "
    "search ranking or buyer conversion, most damaging first. section is one of titles, tags, category, "
    "pricing, description, faq, requirements, gallery. severity is high, medium or low. problem max 140 chars, "
    "fix max 180 chars and must be specific (quote keywords or numbers).\n"
    "STRICT JSON: {\"issues\": [{\"section\": \"titles\", \"severity\": \"high\", \"problem\": \"...\", \"fix\": \"...\"}]}"
)




def _algo_audit(ctx: dict) -> list[dict]:
    """Return basic heuristic audit issues when no LLM is available."""
    kw = ctx["keyword"]
    a = ctx["anchors"]
    original = ctx.get("original") or {}
    issues = []
    title = original.get("title", "")
    if title and kw.lower() not in title.lower():
        issues.append({"section": "titles", "severity": "high",
                        "problem": f"Main keyword '{kw}' missing from gig title — hurts search ranking.",
                        "fix": f"Start your title with '{kw}' or a close variant."})
    desc = original.get("description", "")
    if desc and len(desc) < 600:
        issues.append({"section": "description", "severity": "high",
                        "problem": f"Description only {len(desc)} chars — Fiverr recommends 600-1200 chars.",
                        "fix": "Expand the description: add process steps, FAQ summary, and a CTA."})
    tags = original.get("tags") or []
    if len(tags) < 5:
        issues.append({"section": "tags", "severity": "medium",
                        "problem": f"Only {len(tags)} tags used — Fiverr allows 5.",
                        "fix": "Add tags from top autocomplete phrases: " + ", ".join((ctx.get("autocomplete") or [])[:3])})
    pkgs = original.get("packages") or []
    if pkgs:
        prices = [p.get("price", 0) for p in pkgs if p.get("price")]
        if prices and min(prices) < a["basic"]:
            issues.append({"section": "pricing", "severity": "medium",
                            "problem": f"Basic price ${min(prices)} is below market median ${a['median']}.",
                            "fix": f"Raise Basic to at least ${a['basic']} to align with top sellers."})
    if not pkgs or len(pkgs) < 3:
        issues.append({"section": "pricing", "severity": "high",
                        "problem": "Less than 3 packages — missing Basic/Standard/Premium tiering.",
                        "fix": "Create all three tiers to capture buyers across budget levels."})
    faq = original.get("faq") or []
    if len(faq) < 5:
        issues.append({"section": "faq", "severity": "low",
                        "problem": f"Only {len(faq)} FAQ items — 5-8 recommended to improve conversion.",
                        "fix": "Add FAQs on: revisions, delivery, file formats, customisation, rights."})
    return issues[:8]


async def generate_audit(ctx: dict) -> list[dict]:
    chat = _gig_chat()
    if chat is None:
        # Algorithmic audit when no LLM is available
        return _algo_audit(ctx)
    try:
        data = _parse_json(await run_chat(chat, f"{_ctx_block(ctx)}\n{AUDIT_PROMPT}"), {})
    except Exception as exc:
        logger.warning("LLM audit failed (%s) — using algorithmic audit", exc)
        return _algo_audit(ctx)
    out = []
    for i in data.get("issues") or []:
        if not isinstance(i, dict) or not i.get("problem"):
            continue
        out.append({
            "section": i.get("section") if i.get("section") in SECTIONS else "description",
            "severity": i.get("severity") if i.get("severity") in ("high", "medium", "low") else "medium",
            "problem": str(i["problem"])[:200], "fix": str(i.get("fix", ""))[:260],
        })
    return out[:8]


def original_text(o: dict) -> str:
    parts = [f"Title: I will {o.get('title', '')}", f"Tags: {', '.join(o.get('tags') or [])}",
             f"Description: {o.get('description', '')}"]
    for p in o.get("packages") or []:
        parts.append(f"Package {p.get('tier')}: {p.get('name', '')} — {p.get('description', '')} — "
                     f"${p.get('price', '')}, {p.get('delivery_days', '')} days")
    for f in o.get("faq") or []:
        parts.append(f"FAQ: {f.get('question', '')} / {f.get('answer', '')}")
    return "\n".join(parts)[:3500]


def original_sections(o: dict) -> dict:
    s = {}
    if o.get("title"):
        s["titles"] = [{"text": o["title"]}]
    if o.get("tags"):
        s["tags"] = o["tags"]
    if o.get("description"):
        s["description"] = o["description"]
    pk = [p for p in o.get("packages") or [] if p.get("price")]
    if pk:
        s["pricing"] = {"packages": [{"tier": p.get("tier", ""), "name": p.get("name", ""),
                                      "description": p.get("description", ""), "price": float(p["price"])} for p in pk]}
    if o.get("faq"):
        s["faq"] = o["faq"]
    return s


def before_score(keyword: str, original: dict) -> dict:
    secs = original_sections(original)
    items = build_checklist({"keyword": keyword, "sections": secs}, set())
    items = [c for c in items if c["section"] in secs]
    return {"checklist": items, "passed": sum(c["status"] == "pass" for c in items), "total": len(items)}


PROMPTS = {
    "titles": (
        "Write 3 Fiverr gig title options. Each title is the text AFTER the fixed prefix 'I will' "
        "(do NOT include 'I will'). Max 80 characters each. Ranking-first: a dense keyword string that "
        "still reads sensibly, main keyword first, then comma-separated long-tail variants from autocomplete. "
        "Example: 'design minimalist logo, modern business logo, brand identity logo design'.\n"
        "STRICT JSON: {\"titles\": [\"...\", \"...\", \"...\"]}"
    ),
    "category": (
        "Pick the best matching real Fiverr category, subcategory and service type for this gig.\n"
        "STRICT JSON: {\"category\": \"...\", \"subcategory\": \"...\", \"service_type\": \"...\"}"
    ),
    "tags": (
        "Write 8 Fiverr search tag candidates ranked strongest first, taken from the strongest autocomplete "
        "and proximity phrases. Each tag 20 characters or fewer, lowercase, 1-3 words. Avoid repeating a "
        "single title word exactly unless it is a high-volume phrase.\n"
        "STRICT JSON: {\"tags\": [\"...\"]}"
    ),
    "pricing": (
        "Design three Fiverr packages Basic / Standard / Premium. Use EXACTLY these prices: "
        "Basic ${basic}, Standard ${standard}, Premium ${premium}. Package name max 35 chars "
        "(keyword-rich), package description max 100 chars, delivery_days integer, revisions integer "
        "(or \"Unlimited\" for Premium), included = 2-5 short deliverable items. Also suggest 2-3 gig extras "
        "(e.g. fast delivery) with a USD price and short description.\n"
        "STRICT JSON: {\"packages\": [{\"tier\": \"Basic\", \"name\": \"...\", \"description\": \"...\", "
        "\"delivery_days\": 3, \"revisions\": 1, \"included\": [\"...\"], \"price\": 0}], "
        "\"extras\": [{\"name\": \"...\", \"price\": 0, \"description\": \"...\"}]}"
    ),
    "description": (
        "Write the Fiverr gig description, 900-1150 characters. The MAIN KEYWORD must appear in the first "
        "sentence; weave secondary keywords naturally. Fiverr-safe formatting: short sections with plain-text "
        "headings in CAPS, bullet points using '•', end with a call to action to message or order. Never "
        "mention email, phone, WhatsApp, Skype, PayPal, links or anything off-platform. No markdown.\n"
        "STRICT JSON: {\"description\": \"...\"}"
    ),
    "faq": (
        "Write 6 Fiverr gig FAQ pairs answering common buyer doubts (process, revisions, delivery, "
        "files, customization). Use keywords naturally. Question max 80 chars, answer max 300 chars.\n"
        "STRICT JSON: {\"faq\": [{\"question\": \"...\", \"answer\": \"...\"}]}"
    ),
    "requirements": (
        "Write 4 buyer requirement questions the buyer answers after ordering. type is one of "
        "free_text, multiple_choice, file_upload. multiple_choice needs 2-5 options.\n"
        "STRICT JSON: {\"requirements\": [{\"question\": \"...\", \"type\": \"free_text\", "
        "\"options\": [], \"required\": true}]}"
    ),
    "gallery": (
        "Suggest 3 gallery image file names made of ranking keywords joined by '-' ending in .jpg "
        "(e.g. minimalist-logo-design-modern-business-logo.jpg), each with a short note (max 120 chars) "
        "on what the image should show.\n"
        "STRICT JSON: {\"gallery\": [{\"filename\": \"...\", \"note\": \"...\"}]}"
    ),
}


# ---------------- sanitizers ----------------
def _trim_words(text: str, limit: int) -> str:
    text = re.sub(r"\s+", " ", text or "").strip()
    if len(text) <= limit:
        return text
    cut = text[:limit]
    for sep in (",", " "):
        i = cut.rfind(sep)
        if i > limit * 0.5:
            return cut[:i].rstrip(" ,")
    return cut.rstrip()


def _keyword_pool(ctx: dict) -> list[str]:
    pool = [ctx["keyword"].lower()] + [s.lower() for s in ctx["autocomplete"]] + [t.lower() for t in ctx["tags"]]
    return list(dict.fromkeys(p for p in pool if p))


def clean_titles(data: dict, ctx: dict, keep_order: bool = False) -> list[dict]:
    pool = _keyword_pool(ctx)
    out = []
    for t in (data.get("titles") or [])[:5]:
        t = re.sub(r"^\s*i will\s+", "", str(t), flags=re.I)
        t = _trim_words(t, LIMITS["title"])
        if not t:
            continue
        low = t.lower()
        covered = [p for p in pool if p in low]
        out.append({"text": t, "keywords_covered": covered, "coverage": len(covered)})
    if not keep_order:
        out.sort(key=lambda x: x["coverage"], reverse=True)
    return out[:3]


def clean_tags(data: dict, ctx: dict) -> list[str]:
    cands = list(data.get("tags") or []) + list(ctx["autocomplete"]) + list(ctx["tags"])
    out = []
    for t in cands:
        t = re.sub(r"[^a-z0-9 ]", "", str(t).lower()).strip()
        t = re.sub(r"\s+", " ", t)
        if t and len(t) <= LIMITS["tag"] and t not in out:
            out.append(t)
        if len(out) == LIMITS["tags"]:
            break
    return out


def clean_category(data: dict, ctx: dict) -> dict:
    return {k: str(data.get(k) or "")[:80] for k in ("category", "subcategory", "service_type")}


def clean_pricing(data: dict, ctx: dict) -> dict:
    a = ctx["anchors"]
    prices = {"basic": a["basic"], "standard": a["standard"], "premium": a["premium"]}
    pkgs = []
    for tier in ("Basic", "Standard", "Premium"):
        src = next((p for p in data.get("packages") or [] if str(p.get("tier", "")).lower() == tier.lower()), {})
        rev = src.get("revisions", 1)
        if not (isinstance(rev, int) or str(rev).lower() == "unlimited"):
            rev = 1
        try:
            days = max(1, int(src.get("delivery_days") or 3))
        except (TypeError, ValueError):
            days = 3
        pkgs.append({
            "tier": tier,
            "name": _trim_words(src.get("name") or f"{tier} {ctx['keyword']}", LIMITS["package_name"]),
            "description": _trim_words(src.get("description") or "", LIMITS["package_description"]),
            "delivery_days": days,
            "revisions": "Unlimited" if str(rev).lower() == "unlimited" else rev,
            "included": [str(i)[:60] for i in (src.get("included") or [])][:5],
            "price": prices[tier.lower()],
        })
    extras = []
    for e in (data.get("extras") or [])[:3]:
        try:
            p = max(LIMITS["min_price"], _snap(float(e.get("price") or 10)))
        except (TypeError, ValueError):
            p = 10
        extras.append({"name": str(e.get("name") or "")[:60], "price": p,
                       "description": str(e.get("description") or "")[:120]})
    return {"packages": pkgs, "extras": extras, "anchors": a}


def _strip_off_platform(text: str) -> str:
    text = EMAILISH.sub("", text or "")
    return OFF_PLATFORM.sub("", text)


def clean_description(data: dict, ctx: dict) -> str:
    text = _strip_off_platform(str(data.get("description") or "")).strip()
    text = re.sub(r"[ \t]+", " ", text)
    if len(text) > LIMITS["description"]:
        cut = text[:LIMITS["description"]]
        i = max(cut.rfind("."), cut.rfind("\n"))
        text = cut[:i + 1] if i > 600 else cut
    return text.strip()


def clean_faq(data: dict, ctx: dict) -> list[dict]:
    out = [{"question": _strip_off_platform(str(f.get("question", "")))[:120].strip(),
            "answer": _strip_off_platform(str(f.get("answer", "")))[:400].strip()}
           for f in (data.get("faq") or []) if f.get("question") and f.get("answer")]
    return out[:8]


def clean_requirements(data: dict, ctx: dict) -> list[dict]:
    out = []
    for r in (data.get("requirements") or [])[:5]:
        typ = r.get("type") if r.get("type") in ("free_text", "multiple_choice", "file_upload") else "free_text"
        opts = [str(o)[:60] for o in (r.get("options") or [])][:5] if typ == "multiple_choice" else []
        if r.get("question"):
            out.append({"question": str(r["question"])[:200], "type": typ, "options": opts,
                        "required": bool(r.get("required", True))})
    return out


def _slug(name: str) -> str:
    base = re.sub(r"\.(jpe?g|png|webp)$", "", str(name).lower())
    base = re.sub(r"[^a-z0-9]+", "-", base).strip("-")
    return f"{base[:90].rstrip('-')}.jpg"


def clean_gallery(data: dict, ctx: dict) -> list[dict]:
    return [{"filename": _slug(g.get("filename", "")), "note": str(g.get("note", ""))[:160]}
            for g in (data.get("gallery") or []) if g.get("filename")][:LIMITS["gallery"]]


CLEANERS = {
    "titles": clean_titles, "category": clean_category, "tags": clean_tags, "pricing": clean_pricing,
    "description": clean_description, "faq": clean_faq, "requirements": clean_requirements,
    "gallery": clean_gallery,
}



# ──────────────────────────────────────────────────────────────────────────────
# Algorithmic (no-LLM) fallbacks — used when LLM key absent or call fails
# ──────────────────────────────────────────────────────────────────────────────

def _slug_words(text: str) -> str:
    """Lower-case, hyphen-joined slug from a phrase."""
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def _algo_titles(ctx: dict) -> dict:
    kw = ctx["keyword"]
    pool = _keyword_pool(ctx)
    extras = [p for p in pool if p != kw.lower()][:4]
    combos = [
        f"{kw}, {', '.join(extras[:2])}" if extras else kw,
        f"professional {kw}, {extras[0]}" if extras else f"professional {kw} service",
        f"{kw} for {extras[1] if len(extras) > 1 else 'your business'}, {extras[2] if len(extras) > 2 else kw + ' expert'}",
    ]
    return {"titles": [_trim_words(t.replace("I will ", ""), LIMITS["title"]) for t in combos]}


def _algo_category(ctx: dict) -> dict:
    kw = ctx["keyword"].lower()
    if any(w in kw for w in ("logo", "banner", "design", "graphic", "flyer", "poster", "illustration")):
        return {"category": "Graphics & Design", "subcategory": "Logo Design", "service_type": "Business Logo"}
    if any(w in kw for w in ("write", "blog", "article", "content", "copy", "seo")):
        return {"category": "Writing & Translation", "subcategory": "Articles & Blog Posts", "service_type": "Blog Post"}
    if any(w in kw for w in ("video", "edit", "motion", "animation", "reel")):
        return {"category": "Video & Animation", "subcategory": "Video Editing", "service_type": "YouTube Video"}
    if any(w in kw for w in ("code", "python", "javascript", "react", "web", "app", "api", "bot", "scrape")):
        return {"category": "Programming & Tech", "subcategory": "Web Programming", "service_type": "Full Stack"}
    if any(w in kw for w in ("social media", "instagram", "tiktok", "facebook", "twitter", "marketing")):
        return {"category": "Digital Marketing", "subcategory": "Social Media Marketing", "service_type": "Social Media Management"}
    if any(w in kw for w in ("voice", "audio", "podcast", "music", "sound")):
        return {"category": "Music & Audio", "subcategory": "Voice Over", "service_type": "Business"}
    return {"category": "Programming & Tech", "subcategory": "Web Programming", "service_type": "General"}


def _algo_tags(ctx: dict) -> dict:
    pool = _keyword_pool(ctx)
    tags = []
    for t in pool:
        t = re.sub(r"[^a-z0-9 ]", "", t.lower()).strip()
        t = re.sub(r"\s+", " ", t)
        if t and len(t) <= LIMITS["tag"] and t not in tags:
            tags.append(t)
        if len(tags) == LIMITS["tags"]:
            break
    return {"tags": tags}


def _algo_pricing(ctx: dict) -> dict:
    a = ctx["anchors"]
    kw = ctx["keyword"]
    days = ctx.get("delivery_days") or []
    basic_days = min(days) if days else 3
    std_days = int(statistics.median(days)) if days else 5
    prem_days = max(days) if days else 7
    return {
        "packages": [
            {
                "tier": "Basic",
                "name": _trim_words(f"Basic {kw}", LIMITS["package_name"]),
                "description": _trim_words(f"Starter {kw} package", LIMITS["package_description"]),
                "delivery_days": basic_days,
                "revisions": 1,
                "included": [f"1 {kw} deliverable", "Source files", "Commercial use"],
                "price": a["basic"],
            },
            {
                "tier": "Standard",
                "name": _trim_words(f"Standard {kw}", LIMITS["package_name"]),
                "description": _trim_words(f"Professional {kw} with revisions", LIMITS["package_description"]),
                "delivery_days": std_days,
                "revisions": 3,
                "included": [f"2 {kw} concepts", "Source files", "Commercial use", "Priority support"],
                "price": a["standard"],
            },
            {
                "tier": "Premium",
                "name": _trim_words(f"Premium {kw}", LIMITS["package_name"]),
                "description": _trim_words(f"Top-tier {kw} — full package", LIMITS["package_description"]),
                "delivery_days": prem_days,
                "revisions": "Unlimited",
                "included": [f"3+ {kw} concepts", "Source files", "Commercial use", "Priority support", "Dedicated revisions"],
                "price": a["premium"],
            },
        ],
        "extras": [
            {"name": "Extra-fast delivery", "price": 10, "description": "Delivery in 24 hours"},
            {"name": "Additional revision", "price": 5, "description": "One extra revision round"},
        ],
    }


def _algo_description(ctx: dict) -> dict:
    kw = ctx["keyword"]
    ac = ctx.get("autocomplete") or []
    secondary = ", ".join(ac[:3]) if ac else kw
    a = ctx["anchors"]
    desc = (
        f"Are you looking for expert {kw} services that rank on Fiverr and deliver real results? "
        f"You've come to the right place.\n\n"
        f"WHAT YOU GET\n"
        f"• Professional {kw} tailored to your brand\n"
        f"• Fast delivery with clear communication\n"
        f"• Unlimited revisions on Premium package\n"
        f"• 100% original, commercial-use work\n\n"
        f"WHY CHOOSE ME\n"
        f"• Specialised in {secondary}\n"
        f"• Packages starting from ${a['basic']} — something for every budget\n"
        f"• Satisfaction guaranteed or full revision\n\n"
        f"HOW IT WORKS\n"
        f"1. Place your order and fill in the requirements form\n"
        f"2. I'll start immediately and keep you updated\n"
        f"3. Receive your completed {kw} and request changes if needed\n\n"
        f"Ready to get started? Send me a message or click Order Now!"
    )
    # Strip to limit
    if len(desc) > LIMITS["description"]:
        desc = desc[:LIMITS["description"]]
    return {"description": desc}


def _algo_faq(ctx: dict) -> dict:
    kw = ctx["keyword"]
    a = ctx["anchors"]
    return {"faq": [
        {"question": f"What is included in the Basic {kw} package?",
         "answer": f"The Basic package includes one {kw} deliverable, source files, and commercial use rights, delivered within the stated timeframe."},
        {"question": "How many revisions do I get?",
         "answer": "Basic includes 1 revision, Standard includes 3, and Premium includes unlimited revisions until you are 100% satisfied."},
        {"question": "What information do you need from me?",
         "answer": f"Please provide your business name, preferred style, color preferences, and any examples you like. The more detail, the better the {kw} result."},
        {"question": "How long does delivery take?",
         "answer": f"Delivery times vary by package — from as fast as {min(ctx.get('delivery_days') or [3])} days for Basic up to {max(ctx.get('delivery_days') or [7])} days for Premium."},
        {"question": "Do I own the rights to the final work?",
         "answer": "Yes. Upon completion and payment, you receive full commercial rights to all deliverables."},
        {"question": f"Can you match my existing brand style for {kw}?",
         "answer": "Absolutely. Share your brand guidelines or examples and I will align the work to your existing identity."},
    ]}


def _algo_requirements(ctx: dict) -> dict:
    kw = ctx["keyword"]
    return {"requirements": [
        {"question": f"Please describe your {kw} project in detail (purpose, audience, tone).",
         "type": "free_text", "options": [], "required": True},
        {"question": "What style do you prefer?",
         "type": "multiple_choice",
         "options": ["Modern & Minimal", "Classic & Professional", "Bold & Creative", "Other — describe below"],
         "required": True},
        {"question": "Do you have any reference examples or inspirations?",
         "type": "free_text", "options": [], "required": False},
        {"question": "Please upload any brand assets (logo, fonts, colors) if available.",
         "type": "file_upload", "options": [], "required": False},
    ]}


def _algo_gallery(ctx: dict) -> dict:
    kw = ctx["keyword"]
    slug = _slug_words(kw)
    tags = ctx.get("autocomplete") or ctx.get("tags") or []
    t1 = _slug_words(tags[0]) if tags else slug + "-sample"
    t2 = _slug_words(tags[1]) if len(tags) > 1 else slug + "-example"
    return {"gallery": [
        {"filename": f"{slug}-portfolio-showcase.jpg",
         "note": f"A polished showcase of {kw} work highlighting quality and style."},
        {"filename": f"{slug}-{t1}-example.jpg",
         "note": f"Side-by-side before/after demonstrating the {kw} process and results."},
        {"filename": f"{slug}-{t2}-delivery.jpg",
         "note": f"Final delivered {kw} file previewed in a realistic mockup context."},
    ]}


_ALGO_GENERATORS = {
    "titles":       _algo_titles,
    "category":     _algo_category,
    "tags":         _algo_tags,
    "pricing":      _algo_pricing,
    "description":  _algo_description,
    "faq":          _algo_faq,
    "requirements": _algo_requirements,
    "gallery":      _algo_gallery,
}


def _algo_fallback(section: str, ctx: dict):
    """Return deterministic section content without any LLM call."""
    gen = _ALGO_GENERATORS.get(section)
    if gen is None:
        return None
    try:
        raw = gen(ctx)
        return CLEANERS[section](raw, ctx)
    except Exception as exc:
        logger.warning("Algorithmic fallback failed for %s: %s", section, exc)
        return None


async def generate_section(section: str, ctx: dict):
    """Try LLM first; fall back to algorithmic generation if LLM unavailable."""
    chat = _gig_chat()
    if chat is not None:
        prompt = PROMPTS[section]
        if section == "pricing":
            a = ctx["anchors"]
            prompt = (prompt
                      .replace("${basic}", f"${a['basic']}")
                      .replace("${standard}", f"${a['standard']}")
                      .replace("${premium}", f"${a['premium']}")
                      )
        try:
            text = await run_chat(chat, f"{_ctx_block(ctx)}\n{prompt}")
            result = CLEANERS[section](_parse_json(text, {}), ctx)
            if result:  # LLM gave usable output
                return result
        except Exception as exc:
            logger.warning("LLM failed for section %s (%s) — using algorithmic fallback", section, exc)
    else:
        logger.info("No LLM key available — using algorithmic fallback for section %s", section)
    # Algorithmic fallback
    return _algo_fallback(section, ctx)


# ---------------- checklist ----------------
def build_checklist(setup: dict, locked: set) -> list[dict]:
    pending = set(setup.get("pending_sections") or [])
    s = setup.get("sections") or {}
    kw = setup["keyword"].lower()
    kw_tokens = [t for t in re.findall(r"[a-z0-9]+", kw) if len(t) > 2]
    items = []

    def add(key, label, section, ok):
        items.append({"key": key, "label": label, "section": section,
                      "status": "locked" if section in locked else ("pending" if section in pending else ("pass" if ok else "fail"))})

    titles = s.get("titles") or []
    t0 = (titles[0]["text"].lower() if titles else "")
    add("title_keyword", "Main keyword at the start of the title", "titles",
        bool(t0) and (kw in t0[:45] or all(tok in t0[:50] for tok in kw_tokens)))
    add("title_length", "Every title within 80 characters", "titles",
        bool(titles) and all(len(t["text"]) <= LIMITS["title"] for t in titles))
    tags = s.get("tags") or []
    add("tags_limits", "5 tags, each 20 characters or fewer", "tags",
        len(tags) == LIMITS["tags"] and all(len(t) <= LIMITS["tag"] for t in tags))
    cat = s.get("category") or {}
    add("category", "Category and subcategory chosen", "category", bool(cat.get("category") and cat.get("subcategory")))
    desc = s.get("description") or ""
    first = re.split(r"(?<=[.!?])\s|\n", desc.strip(), maxsplit=1)[0].lower() if desc else ""
    add("desc_keyword", "Main keyword in the first sentence of the description", "description",
        bool(first) and (kw in first or all(tok in first for tok in kw_tokens)))
    add("desc_length", "Description between 600 and 1,200 characters", "description",
        600 <= len(desc) <= LIMITS["description"])
    add("desc_safe", "No contact details or off-platform words", "description",
        bool(desc) and not OFF_PLATFORM.search(desc) and not EMAILISH.search(desc))
    pk = (s.get("pricing") or {}).get("packages") or []
    add("pricing", "3 packages, rising prices, all at least $5", "pricing",
        len(pk) == 3 and pk[0]["price"] >= LIMITS["min_price"] and pk[0]["price"] < pk[1]["price"] < pk[2]["price"]
        and all(len(p["name"]) <= LIMITS["package_name"] and len(p["description"]) <= LIMITS["package_description"] for p in pk))
    add("faq", "5 to 8 FAQ items", "faq", 5 <= len(s.get("faq") or []) <= 8)
    add("requirements", "3 to 5 buyer requirements", "requirements", 3 <= len(s.get("requirements") or []) <= 5)
    gal = s.get("gallery") or []
    add("gallery", "3 keyword-rich gallery file names", "gallery",
        len(gal) == 3 and all(any(tok in g["filename"] for tok in kw_tokens) for g in gal))
    return items


def _iso_age_days(iso: Optional[str]) -> Optional[float]:
    if not iso:
        return None
    try:
        dt = datetime.fromisoformat(iso)
    except ValueError:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return round((datetime.now(timezone.utc) - dt).total_seconds() / 86400, 1)


def present(setup: dict, pro: bool) -> dict:
    locked = set() if pro else set(SECTIONS) - FREE_SECTIONS
    sections = {k: (None if k in locked else (setup.get("sections") or {}).get(k)) for k in SECTIONS}
    checklist = build_checklist(setup, locked)
    scored = [c for c in checklist if c["status"] in ("pass", "fail")]
    age = _iso_age_days(setup.get("scan_created_at"))
    return {
        "id": setup["id"], "keyword": setup["keyword"], "job_id": setup.get("job_id"),
        "kind": setup.get("kind", "build"),
        "audit": {"original": setup.get("original"), "before": setup.get("before"),
                  "issues": setup.get("audit_issues")} if setup.get("kind") == "audit" else None,
        "status": setup.get("status", "complete"), "pending_sections": setup.get("pending_sections", []),
        "sections": sections, "locked_sections": sorted(locked), "is_pro": pro,
        "failed_sections": setup.get("failed_sections", []),
        "checklist": checklist,
        "score": {"passed": sum(c["status"] == "pass" for c in scored), "total": len(scored)},
        "scan_created_at": setup.get("scan_created_at"), "scan_age_days": age,
        "scan_is_stale": age is not None and age > STALE_DAYS,
        "limits": LIMITS, "created_at": setup.get("created_at"), "updated_at": setup.get("updated_at"),
    }


# ---------------- DOCX ----------------
def build_docx(setup: dict) -> bytes:
    if not HAS_DOCX:
        raise HTTPException(
            503,
            "python-docx not installed on this server. DOCX export unavailable. "
            "Run: pip install python-docx"
        )
    s = setup["sections"]
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(f"Fiverr Gig Setup — {setup['keyword']}", 0)

    doc.add_heading("Step 1 · Overview", 1)
    doc.add_heading("Gig title options", 2)
    for i, t in enumerate(s.get("titles") or [], 1):
        doc.add_paragraph(f"Option {i}: I will {t['text']}  ({len(t['text'])}/80)")
    cat = s.get("category") or {}
    doc.add_heading("Category", 2)
    doc.add_paragraph(f"{cat.get('category', '')} › {cat.get('subcategory', '')} › {cat.get('service_type', '')}")
    doc.add_heading("Search tags", 2)
    doc.add_paragraph(", ".join(s.get("tags") or []))

    doc.add_heading("Step 2 · Pricing", 1)
    pricing = s.get("pricing") or {}
    pk = pricing.get("packages") or []
    if pk:
        table = doc.add_table(rows=7, cols=4)
        table.style = "Table Grid"
        rows = ["", "Name", "Description", "Delivery", "Revisions", "Included", "Price"]
        for r, label in enumerate(rows):
            table.cell(r, 0).text = label
        for c, p in enumerate(pk, 1):
            vals = [p["tier"], p["name"], p["description"], f"{p['delivery_days']} days",
                    str(p["revisions"]), "\n".join(p["included"]), f"${p['price']}"]
            for r, v in enumerate(vals):
                table.cell(r, c).text = v
    if pricing.get("extras"):
        doc.add_heading("Gig extras", 2)
        for e in pricing["extras"]:
            doc.add_paragraph(f"{e['name']} — ${e['price']}: {e['description']}", style="List Bullet")

    doc.add_heading("Step 3 · Description & FAQ", 1)
    doc.add_paragraph(s.get("description") or "")
    doc.add_heading("FAQ", 2)
    for f in s.get("faq") or []:
        doc.add_paragraph(f["question"], style="List Number")
        doc.add_paragraph(f["answer"])

    doc.add_heading("Step 4 · Requirements", 1)
    for r in s.get("requirements") or []:
        extra = f" — options: {', '.join(r['options'])}" if r.get("options") else ""
        doc.add_paragraph(f"{r['question']} [{r['type'].replace('_', ' ')}]{extra}", style="List Number")

    doc.add_heading("Step 5 · Gallery", 1)
    for g in s.get("gallery") or []:
        doc.add_paragraph(f"{g['filename']} — {g['note']}", style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------- routes ----------------
def make_gig_routes(users_col, jobs_col, gigs_col):

    async def _user(user):
        doc = await users_col.find_one({"id": user["id"]}, {"_id": 0})
        if not doc:
            raise HTTPException(404, "User not found")
        return doc

    async def _find_job(user_id: str, body: BuildIn) -> dict:
        q = {"user_id": user_id, "status": "complete", "result": {"$ne": None}}
        if body.job_id:
            q["id"] = body.job_id
        elif body.keyword:
            q["keyword"] = {"$regex": f"^{re.escape(body.keyword.strip())}$", "$options": "i"}
        else:
            raise HTTPException(400, "job_id or keyword required")
        docs = await jobs_col.find(q, {"_id": 0}).sort("created_at", -1).limit(1).to_list(1)
        if not docs or not docs[0].get("result"):
            raise HTTPException(404, {"code": "SCAN_REQUIRED", "message": "Scan this keyword first."})
        return docs[0]

    async def _ctx_from_job(job: dict) -> dict:
        r = job["result"]
        gigs = r.get("gigs") or []
        top = sorted(gigs, key=lambda g: g.get("position", 99))[:10]
        try:
            auto = await asyncio.wait_for(fetch_autocomplete(job["keyword"]), timeout=20)
        except asyncio.TimeoutError:
            auto = {"suggestions": []}
        days = [g.get("delivery_days") for g in top if g.get("delivery_days")]
        return {
            "keyword": job["keyword"],
            "top_titles": [g.get("title", "")[:100] for g in top],
            "tags": (r.get("tags") or [])[:20],
            "autocomplete": (auto.get("suggestions") or [])[:12],
            "anchors": price_anchors(gigs, r.get("metrics") or {}),
            "delivery_days": sorted(set(days))[:6],
        }

    async def _owned(sid: str, user_id: str) -> dict:
        doc = await gigs_col.find_one({"id": sid, "user_id": user_id}, {"_id": 0})
        if not doc:
            raise HTTPException(404, "Gig setup not found")
        return doc

    def _assert_unlocked(section: str, pro: bool):
        if section not in SECTIONS:
            raise HTTPException(400, "Unknown section")
        if not pro and section not in FREE_SECTIONS:
            raise HTTPException(402, {"code": "PAYMENT_REQUIRED", "message": "Upgrade to Pro to unlock this section."})

    async def _run_build(sid: str, ctx: dict, wanted: list[str]):
        async def one(name):
            field = "audit_issues" if name == "audit" else f"sections.{name}"
            val = None
            for _ in range(2):
                try:
                    val = await (generate_audit(ctx) if name == "audit" else generate_section(name, ctx))
                except Exception:
                    val = None
                if val:
                    break
            upd = {"$pull": {"pending_sections": name}, "$set": {field: val}}
            if not val:
                upd["$addToSet"] = {"failed_sections": name}
            await gigs_col.update_one({"id": sid}, upd)
        await asyncio.gather(*(one(n) for n in wanted))
        doc = await gigs_col.find_one({"id": sid}, {"_id": 0, "failed_sections": 1})
        status = "failed" if doc and len(doc.get("failed_sections", [])) == len(wanted) else "complete"
        await gigs_col.update_one({"id": sid}, {"$set": {
            "status": status, "updated_at": datetime.now(timezone.utc).isoformat()}})

    @gig_router.post("/build")
    async def build(body: BuildIn, bg: BackgroundTasks, user=Depends(get_current_user)):
        udoc = await _user(user)
        pro = is_pro(udoc)
        job = await _find_job(user["id"], body)
        ctx = await _ctx_from_job(job)
        wanted = SECTIONS if pro else [s for s in SECTIONS if s in FREE_SECTIONS]
        now = datetime.now(timezone.utc).isoformat()
        setup = {"id": str(uuid.uuid4()), "user_id": user["id"], "keyword": job["keyword"],
                 "job_id": job["id"], "scan_created_at": job.get("created_at"), "context": ctx,
                 "sections": {k: None for k in SECTIONS}, "failed_sections": [], "pending_sections": wanted,
                 "status": "building", "created_at": now, "updated_at": now}
        await gigs_col.insert_one(dict(setup))
        bg.add_task(_run_build, setup["id"], ctx, wanted)
        return present(setup, pro)

    @gig_router.post("/audit")
    async def audit(body: AuditIn, bg: BackgroundTasks, user=Depends(get_current_user)):
        udoc = await _user(user)
        pro = is_pro(udoc)
        job = await _find_job(user["id"], BuildIn(job_id=body.job_id))
        original = body.original.model_dump()
        original["title"] = re.sub(r"^\s*i will\s+", "", original["title"], flags=re.I).strip()
        original["tags"] = [t.strip() for t in original["tags"] if t.strip()][:10]
        original["packages"] = [p for p in original["packages"] if p["name"] or p["price"] is not None]
        ctx = await _ctx_from_job(job)
        ctx["original"] = original_text(original)
        wanted = (SECTIONS if pro else [s for s in SECTIONS if s in FREE_SECTIONS]) + ["audit"]
        now = datetime.now(timezone.utc).isoformat()
        setup = {"id": str(uuid.uuid4()), "user_id": user["id"], "keyword": job["keyword"], "kind": "audit",
                 "job_id": job["id"], "scan_created_at": job.get("created_at"), "context": ctx,
                 "original": original, "before": before_score(job["keyword"], original), "audit_issues": None,
                 "sections": {k: None for k in SECTIONS}, "failed_sections": [], "pending_sections": wanted,
                 "status": "building", "created_at": now, "updated_at": now}
        await gigs_col.insert_one(dict(setup))
        bg.add_task(_run_build, setup["id"], ctx, wanted)
        return present(setup, pro)

    @gig_router.get("")
    async def list_setups(user=Depends(get_current_user)):
        docs = await gigs_col.find({"user_id": user["id"]}, {"_id": 0, "context": 0}) \
            .sort("updated_at", -1).limit(100).to_list(100)
        return {"items": [{
            "id": d["id"], "keyword": d["keyword"], "created_at": d["created_at"], "updated_at": d["updated_at"],
            "kind": d.get("kind", "build"),
            "title": ((d.get("sections") or {}).get("titles") or [{}])[0].get("text"),
        } for d in docs]}

    @gig_router.get("/{sid}")
    async def get_setup(sid: str, user=Depends(get_current_user)):
        udoc = await _user(user)
        return present(await _owned(sid, user["id"]), is_pro(udoc))

    @gig_router.post("/{sid}/regenerate")
    async def regenerate(sid: str, body: RegenIn, user=Depends(get_current_user)):
        udoc = await _user(user)
        pro = is_pro(udoc)
        _assert_unlocked(body.section, pro)
        setup = await _owned(sid, user["id"])
        try:
            value = await generate_section(body.section, setup["context"])
        except Exception:
            raise HTTPException(502, "AI generation failed. Try again.")
        if not value:
            raise HTTPException(502, "AI returned an empty section. Try again.")
        now = datetime.now(timezone.utc).isoformat()
        setup["sections"][body.section] = value
        setup["failed_sections"] = [f for f in setup.get("failed_sections", []) if f != body.section]
        setup["updated_at"] = now
        await gigs_col.update_one({"id": sid}, {"$set": {
            f"sections.{body.section}": value, "failed_sections": setup["failed_sections"], "updated_at": now}})
        return present(setup, pro)

    @gig_router.patch("/{sid}")
    async def edit(sid: str, body: PatchIn, user=Depends(get_current_user)):
        udoc = await _user(user)
        pro = is_pro(udoc)
        _assert_unlocked(body.section, pro)
        setup = await _owned(sid, user["id"])
        ctx = setup["context"]
        v = body.value
        if body.section in ("faq", "requirements", "gallery", "pricing", "category") and not (
                isinstance(v, dict) if body.section in ("pricing", "category")
                else isinstance(v, list) and all(isinstance(i, dict) for i in v)):
            raise HTTPException(400, "Invalid value for this section")
        if body.section == "titles":
            value = clean_titles({"titles": [t["text"] if isinstance(t, dict) else t for t in (v or [])]}, ctx,
                                 keep_order=True)
        elif body.section == "pricing":
            value = clean_pricing(v or {}, ctx)
            for p, src in zip(value["packages"], (v or {}).get("packages") or []):
                try:
                    p["price"] = max(LIMITS["min_price"], int(float(src.get("price"))))
                except (TypeError, ValueError):
                    pass
        else:
            value = CLEANERS[body.section]({body.section: v} if body.section != "category" else (v or {}), ctx)
        now = datetime.now(timezone.utc).isoformat()
        setup["sections"][body.section] = value
        setup["updated_at"] = now
        await gigs_col.update_one({"id": sid}, {"$set": {f"sections.{body.section}": value, "updated_at": now}})
        return present(setup, pro)

    @gig_router.delete("/{sid}")
    async def delete(sid: str, user=Depends(get_current_user)):
        res = await gigs_col.delete_one({"id": sid, "user_id": user["id"]})
        if res.deleted_count == 0:
            raise HTTPException(404, "Gig setup not found")
        return {"deleted": True}

    @gig_router.get("/{sid}/docx")
    async def download(sid: str, user=Depends(get_current_user)):
        udoc = await _user(user)
        if not is_pro(udoc):
            raise HTTPException(402, {"code": "PAYMENT_REQUIRED", "message": "Upgrade to Pro to download."})
        setup = await _owned(sid, user["id"])
        name = re.sub(r"[^a-z0-9]+", "_", setup["keyword"].lower()).strip("_")
        return StreamingResponse(
            io.BytesIO(build_docx(setup)),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f'attachment; filename="gig_setup_{name}.docx"'})

    return gig_router
